"""审查报告导出格式转换。"""

from io import BytesIO
import re


def report_to_docx(report: str) -> bytes:
    """将 Markdown 风格审查报告排版为整洁的合同审查 Word 文档。"""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("导出 Word 报告需要安装 python-docx") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal_style = document.styles["Normal"]
    _set_east_asia_font(normal_style, "宋体")
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    for raw_line in report.splitlines():
        line = raw_line.strip()
        if not line or re.fullmatch(r"[-*_]{3,}", line):
            continue

        heading_match = re.match(r"^#{1,6}\s+(.+)$", line)
        if heading_match:
            level = min(line.count("#"), 6)
            paragraph = document.add_heading(
                _sanitize_text(_strip_inline_markdown(heading_match.group(1))),
                level=level,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
            paragraph.paragraph_format.space_after = Pt(8)
            _set_east_asia_font(paragraph.style, "宋体")
            continue

        list_match = re.match(r"^(?:[-*+]|\d+\.)\s+(.+)$", line)
        if list_match:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run(_sanitize_text(_strip_inline_markdown(list_match.group(1))))
            continue

        paragraph = document.add_paragraph(_sanitize_text(_strip_inline_markdown(line)))
        paragraph.paragraph_format.first_line_indent = Inches(0.3)

    _force_simsun_font(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _set_east_asia_font(style_or_run, font_name: str) -> None:
    """同时设置西文和东亚字体，避免中文在 Word 中回退到默认字体。"""
    from docx.oxml.ns import qn

    style_or_run.font.name = font_name
    style_or_run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _strip_inline_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"(```\w*|```|`|\*\*|__|\*|_)", "", text)
    text = re.sub(r"^\s*\|.*\|\s*$", lambda match: match.group(0).replace("|", " "), text)
    return text.strip()


def _sanitize_text(text: str) -> str:
    """移除 Word XML 不允许保存的控制字符。"""
    def is_valid_xml_char(char: str) -> bool:
        codepoint = ord(char)
        return (
            codepoint in (0x9, 0xA, 0xD)
            or 0x20 <= codepoint <= 0xD7FF
            or 0xE000 <= codepoint <= 0xFFFD
            or 0x10000 <= codepoint <= 0x10FFFF
        )

    return "".join(char for char in text if is_valid_xml_char(char))


def _force_simsun_font(document) -> None:
    """统一文档中所有文字 run 的中西文字体为宋体。"""
    from docx.oxml.ns import qn

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "宋体"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "宋体")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "宋体")
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
            run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "宋体")

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "宋体"
                        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "宋体")
                        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "宋体")
                        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
                        run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "宋体")
