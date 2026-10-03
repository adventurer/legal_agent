import unittest
from io import BytesIO
import json

from docx import Document

from web.report_exporter import report_to_docx


class ReportExporterTests(unittest.TestCase):
    def test_exports_markdown_structure_without_damaging_text(self):
        report = r"""# Review #1

## 1. Terms
Value foo_bar and **bold**.

| Field | Value |
| --- | --- |
| Risk | **High** |
| Note | A \| B |

1. First point
2. Second point
"""

        document = Document(BytesIO(report_to_docx(report)))

        self.assertEqual(document.paragraphs[0].style.name, "Heading 1")
        self.assertEqual(document.paragraphs[0].text, "Review #1")
        paragraph_text = " ".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("foo_bar", paragraph_text)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(
            [cell.text for cell in document.tables[0].rows[1].cells],
            ["Risk", "High"],
        )
        self.assertEqual(
            [cell.text for cell in document.tables[0].rows[2].cells],
            ["Note", "A | B"],
        )

        numbered_items = [
            paragraph for paragraph in document.paragraphs
            if paragraph.text in {"First point", "Second point"}
        ]
        self.assertEqual(len(numbered_items), 2)
        self.assertTrue(all(item.style.name == "List Number" for item in numbered_items))

    def test_exports_valid_json_without_interpreting_markdown_in_values(self):
        report = '```json\n{"title":"# literal heading","value":"**literal text**"}\n```'

        document = Document(BytesIO(report_to_docx(report)))
        exported_json = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertEqual(
            json.loads(exported_json),
            {"title": "# literal heading", "value": "**literal text**"},
        )
        self.assertTrue(all(not paragraph.style.name.startswith("Heading") for paragraph in document.paragraphs))
        self.assertTrue(all(paragraph.style.name != "List Bullet" for paragraph in document.paragraphs))


if __name__ == "__main__":
    unittest.main()