#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文件名: services/doc_loader.py
职责:
1. 多源合同解析：支持 Word (.docx)、PDF (.pdf)、图片 (.png, .jpg, .jpeg, .webp, .bmp) 与纯文本 (.txt, .md)
2. 生产级双轨 PDF 提取：优先毫秒级提取矢量文字；若为扫描件/图片页，自动光栅化为高分图片送入 RapidOCR
3. 层次化防截断状态机：按“第一条/一、”等主条款聚合；严禁将 1.1、(1) 等款项截断；自动切分前言与签署盖章页
4. 自带四模态测试数据生成引擎：自动生成 TXT/DOCX/PNG/PDF 并跑通完整闭环测试
"""

import os
import re
import io
import sys
import tempfile
from pathlib import Path
from typing import List, Dict, Any, Optional

# 1. 基础文档处理依赖
try:
    from docx import Document
except ImportError:
    Document = None

try:
    import fitz  # PyMuPDF（PDF 提取与光栅化核心）
except ImportError:
    fitz = None

try:
    import pypdf
except ImportError:
    try:
        import PyPDF2 as pypdf
    except ImportError:
        pypdf = None

# 2. 图片与 OCR 依赖
try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    Image, ImageDraw, ImageFont = None, None, None

try:
    from rapidocr_onnxruntime import RapidOCR
    _rapid_ocr_engine = RapidOCR()
except ImportError:
    _rapid_ocr_engine = None

try:
    import pytesseract
except ImportError:
    pytesseract = None


class ContractClause:
    """单个合同条款结构体"""
    def __init__(self, index: int, title: str, content: str, raw_text: str, clause_type: str = "article"):
        self.index = index
        self.title = title
        self.content = content
        self.raw_text = raw_text
        self.clause_type = clause_type  # preamble (前言), article (正文条款), sign_page (盖章页)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "index": self.index,
            "title": self.title,
            "content": self.content,
            "raw_text": self.raw_text,
            "type": self.clause_type
        }

    def __repr__(self) -> str:
        return f"<Clause {self.index} [{self.clause_type}]: {self.title[:18]}... ({len(self.content)} 字)>"


class DocumentLoader:
    """合同文档多模态解析与切分器 (工业增强版)"""

    # 一级大条款正则：严格锚定一级条款（如“第一条 违约责任”、“一、 合作事项”）
    MAJOR_CLAUSE_PATTERN = re.compile(
        r"^(?:第[一二三四五六七八九十百千万\d]+条|\b[一二三四五六七八九十百]+[、. ])\s*(.*)$"
    )

    # 签署盖章页识别正则
    SIGN_PAGE_PATTERN = re.compile(
        r"^(?:（以下无正文|以下无正文|双方签署|协议签署盖章页|甲方（盖章）|乙方（盖章）|甲方\(盖章\)|乙方\(盖章\))"
    )

    SUPPORTED_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}

    def __init__(self):
        pass

    def extract_text_from_docx(self, file_path: Path) -> str:
        """从 Word 文档中提取段落与表格文本"""
        if Document is None:
            raise ImportError("未安装 python-docx，请执行: pip install python-docx")

        doc = Document(str(file_path))
        lines = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    unique_texts = list(dict.fromkeys(row_texts))
                    lines.append(" | ".join(unique_texts))

        return "\n".join(lines)

    def extract_text_from_pdf(self, file_path: Path) -> str:
        """
        高可用双轨 PDF 提取方案：
        1. 优先使用 PyMuPDF (fitz) 提取矢量文字。
        2. 扫描件/图片页判定：若单页有效文字少于 15 字符，光栅化为 200 DPI 图片，送入 OCR 兜底。
        3. 次选 pypdf 降级运行。
        """
        lines = []
        path_str = str(file_path)

        # 方案 A: PyMuPDF (fitz)
        if fitz is not None:
            try:
                doc = fitz.open(path_str)
                for page_idx, page in enumerate(doc):
                    page_text = page.get_text("text").strip()
                    # 统计中文字符与英文字符数，过滤掉仅有空白页或页码干扰
                    valid_chars = re.findall(r'[\u4e00-\u9fa5\w]', page_text)
                    
                    if len(valid_chars) >= 15:
                        for line in page_text.splitlines():
                            if line.strip():
                                lines.append(line.strip())
                    else:
                        # 触发 OCR 兜底：纯图片扫描页或复杂印章页
                        print(f"[*] PDF 第 {page_idx+1} 页文字稀少（疑似扫描件），启动 OCR 光栅化识别...", flush=True)
                        pix = page.get_pixmap(dpi=200)
                        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        ocr_text = self._do_ocr(img)
                        for line in ocr_text.splitlines():
                            if line.strip():
                                lines.append(line.strip())
                if lines:
                    return "\n".join(lines)
            except Exception as e:
                print(f"[警告] PyMuPDF 解析失败: {e}，尝试降级引擎...", flush=True)

        # 方案 B: 基础 pypdf 降级
        if pypdf is not None:
            try:
                reader = pypdf.PdfReader(path_str)
                for page in reader.pages:
                    text = page.extract_text() or ""
                    for line in text.splitlines():
                        if line.strip():
                            lines.append(line.strip())
                if lines:
                    return "\n".join(lines)
            except Exception as e:
                print(f"[警告] pypdf 解析失败: {e}", flush=True)

        if not lines:
            raise RuntimeError(
                f"PDF 文件解析失败或为空: {file_path.name}。\n"
                f"请确保安装了 PyMuPDF 与 RapidOCR:\n"
                f"pip install pymupdf rapidocr-onnxruntime pillow"
            )

        return "\n".join(lines)

    def _do_ocr(self, img_obj: Any) -> str:
        """统一底层 OCR 执行器，内置防报错机制"""
        if _rapid_ocr_engine is not None:
            try:
                # 兼容无需显式安装 numpy 的安全转换
                if hasattr(img_obj, "convert"):
                    img_rgb = img_obj.convert("RGB")
                    with io.BytesIO() as buf:
                        img_rgb.save(buf, format="PNG")
                        img_bytes = buf.getvalue()
                    result, _ = _rapid_ocr_engine(img_bytes)
                else:
                    result, _ = _rapid_ocr_engine(img_obj)

                if result:
                    return "\n".join([item[1].strip() for item in result if item[1].strip()])
                return ""
            except Exception as e:
                print(f"[警告] RapidOCR 执行异常: {e}，尝试备选 OCR 引擎...", flush=True)

        if pytesseract is not None:
            try:
                return pytesseract.image_to_string(img_obj, lang="chi_sim+eng")
            except Exception as e:
                print(f"[警告] pytesseract 识别失败: {e}", flush=True)

        return ""

    def extract_text_from_image(self, file_path: Path) -> str:
        print(file_path)
        """从原生图片文件中提取文字"""
        if Image is None:
            raise ImportError("未安装 Pillow 库，请执行: pip install pillow")
        img = Image.open(file_path)
        text = self._do_ocr(img)
        if not text.strip():
            raise RuntimeError(f"未能从图片中提取到有效文本: {file_path.name}")
        return text

    def extract_text(self, file_path: str | Path) -> str:
        """通用格式路由分发"""
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {path}")

        suffix = path.suffix.lower()

        if suffix == ".docx":
            raw_text = self.extract_text_from_docx(path)
        elif suffix == ".pdf":
            raw_text = self.extract_text_from_pdf(path)
        elif suffix in self.SUPPORTED_IMAGE_EXTS:
            print(f"[*] 检测到图片格式合同 ({suffix})，正在启动 OCR 文本识别...", flush=True)
            raw_text = self.extract_text_from_image(path)
        elif suffix in [".txt", ".md"]:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                raw_text = f.read()
        else:
            supported = [".docx", ".pdf", ".txt"] + list(self.SUPPORTED_IMAGE_EXTS)
            raise ValueError(f"不支持的文件格式: {suffix}，当前支持格式: {supported}")

        return self._normalize_text(raw_text)

    def _normalize_text(self, text: str) -> str:
        """清洗排版乱码与冗余字符"""
        if not text:
            return ""
        text = text.replace("\t", " ").replace("\u3000", " ")
        lines = [line.strip() for line in text.splitlines()]
        return "\n".join([line for line in lines if line])

    def split_into_clauses(self, full_text: str) -> List[ContractClause]:
        """
        层次化状态机切片：
        1. 严格以“第一条/一、”为大条款边界；
        2. 1.1、1.2、(1)、(2) 等子项全部合并保留在所属大条款正文中，防语义碎裂；
        3. 完整识别并隔离前言（PREAMBLE）与签署盖章页（SIGN_PAGE）。
        """
        lines = [line.strip() for line in full_text.splitlines() if line.strip()]
        clauses: List[ContractClause] = []

        state = "PREAMBLE"  # PREAMBLE -> BODY -> SIGN_PAGE
        preamble_lines = []
        current_title = ""
        current_lines = []
        sign_page_lines = []
        counter = 0

        for line in lines:
            # 1. 盖章签署页判定
            if self.SIGN_PAGE_PATTERN.search(line):
                state = "SIGN_PAGE"
                sign_page_lines.append(line)
                continue

            if state == "SIGN_PAGE":
                sign_page_lines.append(line)
                continue

            # 2. 一级主条款判定
            major_match = self.MAJOR_CLAUSE_PATTERN.match(line)
            if major_match:
                if state == "PREAMBLE":
                    if preamble_lines:
                        counter += 1
                        clauses.append(ContractClause(
                            index=counter,
                            title="合同前言与主体信息",
                            content="\n".join(preamble_lines).strip(),
                            raw_text="\n".join(preamble_lines).strip(),
                            clause_type="preamble"
                        ))
                    state = "BODY"

                elif state == "BODY" and current_title:
                    counter += 1
                    body_content = "\n".join(current_lines).strip()
                    clauses.append(ContractClause(
                        index=counter,
                        title=current_title,
                        content=body_content,
                        raw_text=f"{current_title}\n{body_content}",
                        clause_type="article"
                    ))
                    current_lines = []

                current_title = line
            else:
                # 3. 子条款合并处理
                if state == "PREAMBLE":
                    preamble_lines.append(line)
                elif state == "BODY":
                    # 核心：小款（如 1.1、2.2）或 (1)、(2) 强制合并在当前大条款内
                    current_lines.append(line)

        # 归档最后一个正文条款
        if current_title and current_lines:
            counter += 1
            body_content = "\n".join(current_lines).strip()
            clauses.append(ContractClause(
                index=counter,
                title=current_title,
                content=body_content,
                raw_text=f"{current_title}\n{body_content}",
                clause_type="article"
            ))

        # 归档盖章页
        if sign_page_lines:
            counter += 1
            sign_content = "\n".join(sign_page_lines).strip()
            clauses.append(ContractClause(
                index=counter,
                title="协议签署与盖章页",
                content=sign_content,
                raw_text=sign_content,
                clause_type="sign_page"
            ))

        # 全文未分条款时的降级兜底
        if not clauses and full_text.strip():
            clauses.append(ContractClause(
                index=1,
                title="合同正文",
                content=full_text.strip(),
                raw_text=full_text.strip(),
                clause_type="article"
            ))

        return clauses

    def load_and_split(self, file_path: str | Path) -> List[ContractClause]:
        """对外主接口"""
        text = self.extract_text(file_path)
        return self.split_into_clauses(text)


# ==================== 测试套件：动态生成测试文件引擎 ====================

class MockContractGenerator:
    """动态生成四种类型合同测试用例（TXT, DOCX, PNG, PDF）"""

    SAMPLE_CONTRACT_TEXT = (
        "高端智能制造设备采购与维保服务协议\n\n"
        "合同编号：HT-202609-XYZ8890\n"
        "甲方：北京华创智能科技股份有限公司\n"
        "乙方：苏州精工自动化系统工程有限公司\n\n"
        "第一条 交付期限与违约金\n"
        "1.1 设备明细：乙方须向甲方提供6轴工业机器人12台。\n"
        "1.2 逾期交付：乙方逾期交付的，每日应按合同总金额的 5% 向甲方支付违约金。\n\n"
        "第二条 货款结算与付款周期\n"
        "2.1 验收标准：经连续 72 小时无故障运行后签署最终合格单。\n"
        "2.2 货款分期结算：\n"
        "（1）首付款：生效后支付 15%；\n"
        "（2）尾款：视甲方内部季度资金周转充裕情况安排支付，不受商业惯例限制。\n\n"
        "第三条 争议解决与独任仲裁\n"
        "3.1 因本合同引起的一切争议，由甲方指定的单方独任仲裁员在其个人办公场所内裁决，裁决为终局。\n\n"
        "（以下无正文，为协议签署盖章页）\n"
        "甲方（盖章）：北京华创智能科技股份有限公司\n"
        "乙方（盖章）：苏州精工自动化系统工程有限公司"
    )

    @classmethod
    def generate_all(cls, output_dir: Path) -> Dict[str, Path]:
        output_dir.mkdir(parents=True, exist_ok=True)
        paths = {}

        # 1. 生成纯文本文件 (.txt)
        txt_path = output_dir / "test_contract.txt"
        txt_path.write_text(cls.SAMPLE_CONTRACT_TEXT, encoding="utf-8")
        paths["txt"] = txt_path

        # 2. 生成 Word 文件 (.docx)
        if Document is not None:
            docx_path = output_dir / "test_contract.docx"
            doc = Document()
            for line in cls.SAMPLE_CONTRACT_TEXT.split("\n"):
                if line.strip():
                    doc.add_paragraph(line.strip())
            doc.save(str(docx_path))
            paths["docx"] = docx_path

        # 3. 生成图片文件 (.png)
        if Image is not None and ImageDraw is not None:
            img_path = output_dir / "test_contract.png"
            img = Image.new("RGB", (950, 750), color=(255, 255, 255))
            draw = ImageDraw.Draw(img)

            font = None
            font_candidates = [
                "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
                "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                "simsun.ttc",
                "msyh.ttc",
            ]
            for font_file in font_candidates:
                if os.path.exists(font_file):
                    try:
                        font = ImageFont.truetype(font_file, 16)
                        break
                    except Exception:
                        continue
            if font is None:
                font = ImageFont.load_default()

            y = 25
            for line in cls.SAMPLE_CONTRACT_TEXT.split("\n"):
                if line.strip():
                    draw.text((35, y), line.strip(), fill=(0, 0, 0), font=font)
                    y += 26
                else:
                    y += 10

            img.save(str(img_path))
            paths["image"] = img_path

        # 4. 生成原生 PDF 文件 (.pdf)
        pdf_path = output_dir / "test_contract.pdf"
        if fitz is not None:
            # 优先用 PyMuPDF 直接构建测试 PDF，免除额外三方库依赖
            doc_pdf = fitz.open()
            page = doc_pdf.new_page(width=595, height=842)  # A4 页面
            text_point = fitz.Point(40, 50)
            page.insert_text(text_point, cls.SAMPLE_CONTRACT_TEXT, fontsize=11)
            doc_pdf.save(str(pdf_path))
            doc_pdf.close()
            paths["pdf"] = pdf_path
        else:
            try:
                from reportlab.pdfgen import canvas
                from reportlab.lib.pagesizes import letter
                c = canvas.Canvas(str(pdf_path), pagesize=letter)
                y = 750
                for line in cls.SAMPLE_CONTRACT_TEXT.split("\n"):
                    if line.strip():
                        c.drawString(40, y, line.strip())
                        y -= 22
                c.save()
                paths["pdf"] = pdf_path
            except ImportError:
                pass

        return paths


# ==================== 自动化全闭环测试 ====================
if __name__ == "__main__":
    print("=" * 70)
    print(" 🚀 DocumentLoader 增强版自检：多模态合同提取与层次化切片测试")
    print("=" * 70)

    test_dir = Path(tempfile.gettempdir()) / "legal_doc_loader_test_v3"
    print(f"[*] 正在自动生成测试套件到临时目录: {test_dir}")
    generated_files = MockContractGenerator.generate_all(test_dir)

    for f_type, f_path in generated_files.items():
        print(f"  - [{f_type.upper()}] 成功构建: {f_path.name}")

    print("\n" + "=" * 70)
    print("[*] 开始执行各文件格式解析与条款切片验证:")
    loader = DocumentLoader()

    for f_type, f_path in generated_files.items():
        print(f"\n>>> 正在验证 [{f_type.upper()}] 格式: {f_path.name}")
        try:
            clauses = loader.load_and_split(f_path)
            print(f"    [成功] 解析出 {len(clauses)} 个语义结构单元")
            for c in clauses:
                print(f"     * [{c.clause_type.upper():<9}] 序号 {c.index}: {c.title:<22} (正文 {len(c.content)} 字符)")
        except Exception as e:
            print(f"    [失败] 发生异常: {e}")

    print("\n" + "=" * 70)
    print("全部测试流程执行完毕！")